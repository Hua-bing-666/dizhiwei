"""Revise the user's current Word file without rebuilding the school form."""
from pathlib import Path
from copy import deepcopy
import hashlib, json, re
from docx import Document
from docx.text.paragraph import Paragraph
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'tmp/table_reference_review'
OUT.mkdir(parents=True, exist_ok=True)
SOURCE = ROOT / 'deliverables/大创申报书_格式优化稿.docx'
DEST = ROOT / 'deliverables/大创申报书_表格与文献修订稿.docx'
digest = lambda p: hashlib.sha256(p.read_bytes()).hexdigest()
original_hash = digest(SOURCE)
(OUT / 'user_input.docx').write_bytes(SOURCE.read_bytes())
source_md = ROOT / 'application/正式模板正文.md'
frozen_dir=ROOT/'application/table_reference_revision'
frozen_dir.mkdir(parents=True,exist_ok=True)
frozen_original=frozen_dir/'original_body.md'
old_md = source_md.read_text(encoding='utf-8')
if '表格与文献修订申报正文' in old_md:
    old_md = (frozen_original if frozen_original.exists() else OUT/'original_body.md').read_text(encoding='utf-8')
if not frozen_original.exists():frozen_original.write_text(old_md,encoding='utf-8',newline='\n')
(OUT / 'original_body.md').write_text(old_md, encoding='utf-8')
library = json.loads((ROOT / 'application/references.json').read_text(encoding='utf-8'))
numbers = [1, 2, 7, 8, 10, 12, 14, 15]
selected = []
for new, old in enumerate(numbers, 1):
    ref = deepcopy(next(r for r in library['references'] if r['number'] == old))
    ref['library_number'], ref['number'] = old, new
    ref['citation'] = re.sub(r'^\[\d+\]', f'[{new}]', ref['citation'])
    selected.append(ref)
assert [r['language'] for r in selected].count('中文') == 5
selection = {'date': '2026-10-05', 'scope': 'current_application_only', 'chinese': 5,
             'english': 3, 'references': selected, 'research_library_preserved': True}
(ROOT / 'application/申报精选参考文献.json').write_text(json.dumps(selection, ensure_ascii=False, indent=2), encoding='utf-8',newline='\n')

paragraphs = [
('1 南疆应用需求与研究对象',
'李慧琴等结合南疆棉花“滴水出苗”提出压力均匀、铺管及用水管理建议[1]，说明运行压力是区域滴灌管理中需要关注的条件。这一研究支持选择南疆滴灌巡检场景，但不能据此推定本原型具有节水或增产效果。本期以支管水力故障的可观测变化为对象，实际轮灌方式、安装条件与维护需求仍需现场核对。'),
('2 滴灌水力模型与测量依据',
'丁法龙等从变质量流动分析滴灌管沿程压力，并通过测压、测流试验验证模型[2]。实际滴灌管的压力分布需要考虑分布出流及结构条件，本项目的集中平方律模型仅用于监测区段的初步设计，不能替代分布模型和实测标定。实施时先检验健康压力—流量关系及残差，再冻结特征和阈值；以布点、两通道同步误差和差压不确定度的联合评价确定低成本测量的适用范围。'),
('3 间接监测与早期评价边界',
'Petit等综述滴灌堵塞的检测与分析方法，讨论压力、流量等间接监测方法及其适用限制[6]。水力变化不能直接说明堵塞材料、形成机理或单滴头的微观堵塞起点。因此本项目采用清水节流阀验证等效受限的水力响应；“早期”按独立参考测量确认的故障程度和预设失效参照评价。若微小变化被测量不确定度覆盖，应报告可辨识边界，不能用平滑曲线代替检测证据。'),
('4 多工况诊断与可控验证',
'付建民等通过液相管道试验与模拟研究压力、流量对小孔泄漏速率的影响[3]；龚骏等研究多工况下管道泄漏检测[4]。这些工作提示应同时分析水力耦合和正常调节干扰，但其管道参数与识别率不能直接移植到滴灌。Loukatos等利用可控灌溉平台组织压力、流量和无线数据采集[7]，说明可控台架及物联网集成已有研究基础。本项目拟以独立操作标签、同信息强基线和跨批次留出验证研究渐进故障的辨识条件，保留未知正常负载、压力漂移及切换期间的失败实例。'),
('5 农业物联网与通信评价',
'孟祥等完成基于LoRa技术的农田智能灌溉系统设计[5]，为传感节点、通信和管理软件的组织提供国内研究实例。本项目将系统交付细化为测量节点、接收存储、有限诊断和维护复测，不能仅以采用LoRa或增加页面作为创新。Lozoya等研究自触发灌溉控制中的无线通信与能量取舍，其实际链路采用IEEE 802.15.4/Zigbee，并非LoRa[8]。本项目借鉴其对照评价思路，在相同检测器、载荷和确认条件下比较周期与事件策略，分别测量本地检测延迟、远端持久化延迟和整机能量，不移用该文的性能结果，也不以发送帧数推算真实续航。'),
('6 本项目拟研究的增量',
'上述五篇中文文献构成区域应用、水力分析、泄漏诊断和LoRa系统组织的主要依据，三篇英文文献补充间接监测边界、可控实验平台与通信评价。拟研究的增量集中于两个问题：低成本双压力流量测量与运行信息在何种故障程度和误差条件下能够改善巡检提示；在统一检测与确认机制下，减少上报是否仍能维持远端及时告警。通过可重复台架、独立真值、跨批次留出与失败案例回答问题，避免将系统集成规模或仿真图数量当作科研贡献。')]
start = old_md.index('## 四、'); end = old_md.index('## 五、')
policy = next(s for s in old_md[start:end].splitlines() if s.startswith('制度参考：'))
new_literature = '## 四、国、内外研究现状和发展动态\n\n' + '\n\n'.join('### '+h+'\n\n'+b for h,b in paragraphs)
new_literature += '\n\n参考文献\n\n' + '\n\n'.join(r['citation'] for r in selected) + '\n\n' + policy + '\n\n'
replacements = {
'不替代分布出流管网或CFD标定[2,3,12]。':'不替代分布出流管网的实测标定[2,6]。',
'结项后优先研究跨工况迁移：增加负载档和日期，比较冻结参照、少量健康重标定与经验模型，整档留出评价误差和适应成本[2,12,13]。新增节点、公共供压和同步日志后，再研究多支管共同扰动、局部异常与事件竞争[10,13]。长期维护配对记录可用于检验参照漂移、复测稳定性及核查工作量[12,16]。扩展须落实数据和资源，基础任务仍为单区段验证。':
'结项后优先研究跨工况迁移：增加负载档和日期，比较冻结参照、少量健康重标定与经验模型，整档留出评价适应成本；扩展依据为水力关系与间接监测边界[2,6]。新增节点及同步日志后，拟研究多支管共同扰动与事件竞争；积累维护配对记录后，拟检验参照漂移和复测稳定性。两项扩展均以资源和数据落实为条件，不计入当前基础交付。',
'已形成17篇文献及正文对应，其中中文11篇、英文6篇；':'已完成8篇申报参考文献的筛选与正文对应，其中中文5篇、英文3篇；',
'前期模型与软件由辅助工具协助形成，作为可复现的方案基础；申报与实施还需成员能够独立解释、运行和修改核心代码。':'前期已形成可复现的模型与软件方案，申报与实施要求成员能够独立解释、运行和修改核心代码。'
,'运行Web已接入当前版本24条冻结带噪观测，K/R̂取模型10样本均值。本轮六页面四宽度及横屏、20项软件流程和3个实际导出检查通过，截图来自原生浏览器；事件冻结来源与模型版本，离线/无效禁止恢复归档，人工操作与ACK分开。当前Web没有接入主机SQLite，独立主机检查不能相加称系统端到端完成。':
'运行Web可回放当前版本的24条带噪观测，K/R̂采用10样本均值，已完成六模块页面、事件处置和JSON/CSV导出的浏览器检查。事件记录保留数据来源与模型版本，离线或无效状态下禁止恢复归档，人工操作与无线ACK分别记录。当前Web尚未接入主机SQLite，主机事务检查不作为系统联调完成的证明。'
}
new_md = old_md[:start] + new_literature + old_md[end:]
for a,b in replacements.items():
    assert a in new_md, a[:40]
    new_md = new_md.replace(a,b)
new_md = new_md.replace('2026年10月5日格式优化申报正文。','2026年10月5日表格与文献修订申报正文。')
source_md.write_text(new_md, encoding='utf-8', newline='\n')

doc = Document(SOURCE)
def replace_text(p, value, heading=None):
    if heading is None and re.match(r'^\d+(\.\d+)*\s', value) and '\n' in value:
        heading,value=value.split('\n',1)
    p.clear()
    if heading is None: p.add_run(value)
    else:
        for line in heading.split('\n'):
            p.add_run(line).bold = True
            p.add_run().add_break()
        p.add_run(value)
    for run in p.runs:
        run.font.name = '宋体'; run.font.size = Pt(10.5)
        run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '宋体')

for node in doc.element.xpath('.//w:p'):
    p = Paragraph(node, doc)
    value = p.text
    for a,b in replacements.items(): value = value.replace(a,b)
    if value != p.text:
        replace_text(p, value)
        p.paragraph_format.widow_control=True

lit_cell = next(c for c in doc.element.xpath('.//w:tc') if ''.join(c.xpath('.//w:t/text()')).startswith('四、国、内外研究现状和发展动态'))
holders = lit_cell.findall(qn('w:tbl'))
assert len(holders) == 7
for table,(heading,body) in zip(holders[:6],paragraphs):
    nodes = table.xpath('.//w:p')
    target = next(n for n in nodes if ''.join(n.xpath('.//w:t/text()')))
    replace_text(Paragraph(target, doc), body, heading)
bibliography = holders[6]
rows = bibliography.findall(qn('w:tr'))
assert len(rows) == 17
for row in rows: bibliography.remove(row)
for old,ref in zip(numbers,selected):
    row = rows[old-1]
    target = row.xpath('.//w:p')[0]
    replace_text(Paragraph(target, doc), ref['citation'])
    bibliography.append(row)

# Keep school headings with their first body block in WPS, including across rows.
for label in ['四、国、内外研究现状和发展动态', '六、项目预期成果', '八、已有基础']:
    cell=next(c for c in doc.element.xpath('.//w:tc') if ''.join(c.xpath('.//w:t/text()')).startswith(label))
    title=next(p for p in cell.findall(qn('w:p')) if ''.join(p.xpath('.//w:t/text()'))==label)
    next_node=title.getnext()
    target=next_node if next_node.tag==qn('w:p') else next(p for p in next_node.xpath('.//w:p') if ''.join(p.xpath('.//w:t/text()')))
    p=Paragraph(target,doc);value=p.text
    parts=value.split('\n',1)
    if len(parts)==2 and re.match(r'^\d+(\.\d+)*\s',parts[0]):
        replace_text(p,parts[1],label+'\n'+parts[0])
    else:replace_text(p,value,label)
    p.paragraph_format.keep_together=True
    p.paragraph_format.page_break_before=False
    cell.remove(title)
    if next_node.tag==qn('w:p'):
        container=deepcopy(holders[0])
        container_cell=container.findall(qn('w:tr'))[0].findall(qn('w:tc'))[0]
        for child in list(container_cell):
            if child.tag!=qn('w:tcPr'):container_cell.remove(child)
        container_cell.append(deepcopy(target))
        cell.replace(target,container)
# WPS can split this direct outer-cell paragraph despite KeepLines.
for node in doc.element.xpath('.//w:p'):
    if ''.join(node.xpath('.//w:t/text()')).startswith('六场景、四压力随机噪声档'):
        parent=node.getparent()
        container=deepcopy(holders[0])
        cell=container.findall(qn('w:tr'))[0].findall(qn('w:tc'))[0]
        for child in list(cell):
            if child.tag!=qn('w:tcPr'):cell.remove(child)
        cell.append(deepcopy(node));parent.replace(node,container)
reference_title=next(p for p in lit_cell.findall(qn('w:p')) if ''.join(p.xpath('.//w:t/text()'))=='参考文献')
replace_text(Paragraph(bibliography.findall(qn('w:tr'))[0].xpath('.//w:p')[0],doc),selected[0]['citation'],'参考文献')
lit_cell.remove(reference_title)

def element(parent, name, value=None):
    node = parent.find(qn('w:'+name))
    if node is None: node=OxmlElement('w:'+name); parent.append(node)
    if value is not None: node.set(qn('w:val'), str(value))
    return node
def align(p, alignment):
    p.alignment = alignment
    fmt=p.paragraph_format
    fmt.left_indent=Pt(0);fmt.right_indent=Pt(0);fmt.first_line_indent=Pt(0)
    ind=p._p.get_or_add_pPr().find(qn('w:ind'))
    if ind is not None:
        for attr in ['leftChars','rightChars','firstLineChars','hanging','hangingChars','start','end','startChars','endChars']:
            ind.attrib.pop(qn('w:'+attr),None)
records=[]
for table in doc.element.xpath('.//w:tbl'):
    rows=table.findall(qn('w:tr'))
    if not rows: continue
    headers=[''.join(c.xpath('.//w:t/text()')) for c in rows[0].findall(qn('w:tc'))]
    if headers[0] not in ['工作包','层次','时段','用途']: continue
    pr=table.find(qn('w:tblPr'))
    # Table-property XML order matters to WPS for nested-table centering.
    old_jc=pr.find(qn('w:jc'))
    if old_jc is not None:pr.remove(old_jc)
    pr.get_or_add_jc().val=WD_ALIGN_PARAGRAPH.CENTER
    # WPS ignores nested-table alignment/insets; fill the available cell width.
    # The table and centered caption then share the same actual center in both apps.
    parent_width=int(table.getparent().find(qn('w:tcPr')).find(qn('w:tcW')).get(qn('w:w')))
    wrapper=table.getparent().getparent().getparent()
    if wrapper.tag==qn('w:tbl') and len(wrapper.findall(qn('w:tr')))==1:
        wrapper_pr=wrapper.find(qn('w:tblPr'))
        wrapper_pr.find(qn('w:tblW')).set(qn('w:w'),str(parent_width))
        wrapper_pr.find(qn('w:tblInd')).set(qn('w:w'),'0')
        wrapper.find(qn('w:tblGrid')).find(qn('w:gridCol')).set(qn('w:w'),str(parent_width))
    table_width=int(pr.find(qn('w:tblW')).get(qn('w:w')))
    grid=table.find(qn('w:tblGrid')).findall(qn('w:gridCol'))
    widths=[round(int(c.get(qn('w:w')))*parent_width/table_width) for c in grid]
    widths[-1]=parent_width-sum(widths[:-1])
    for col,width in zip(grid,widths):col.set(qn('w:w'),str(width))
    pr.find(qn('w:tblW')).set(qn('w:w'),str(parent_width))
    for row in rows:
        for cell,width in zip(row.findall(qn('w:tc')),widths):
            cell.find(qn('w:tcPr')).find(qn('w:tcW')).set(qn('w:w'),str(width))
    ind=element(pr,'tblInd');ind.set(qn('w:w'),'0');ind.set(qn('w:type'),'dxa')
    layout=element(pr,'tblLayout');layout.set(qn('w:type'),'fixed');layout.attrib.pop(qn('w:val'),None)
    for ri,row in enumerate(rows):
        rp=row.find(qn('w:trPr'))
        if rp is None:rp=OxmlElement('w:trPr');row.insert(0,rp)
        element(rp,'cantSplit')
        if ri==0:element(rp,'tblHeader','true')
        for ci,cell in enumerate(row.findall(qn('w:tc'))):
            cp=cell.find(qn('w:tcPr'));element(cp,'vAlign','center')
            center = ri==0 or (headers[0]=='用途' and ci==1)
            for n in cell.findall(qn('w:p')):
                p=Paragraph(n,doc);align(p, WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT)
                if ri==0:
                    for run in p.runs:run.bold=True
    records.append({'headers':headers,'rows':len(rows),'alignment':'center','width_pt':parent_width/20,'column_widths_pt':[x/20 for x in widths]})
assert len(records)==4
captions=[]
for n in doc.element.xpath('.//w:p'):
    p=Paragraph(n,doc)
    if re.match(r'^表[1-4]\s',p.text):
        align(p,WD_ALIGN_PARAGRAPH.CENTER)
        p.paragraph_format.keep_with_next=True
        p.paragraph_format.space_before=Pt(6);p.paragraph_format.space_after=Pt(3)
        for run in p.runs:run.bold=True;run.font.size=Pt(10.5)
        captions.append(p.text)
assert len(captions)==4
doc.save(DEST)
assert digest(SOURCE)==original_hash, 'User input changed unexpectedly'
(OUT/'build_record.json').write_text(json.dumps({'input_path':str(SOURCE),'input_sha256':original_hash,
 'output_path':str(DEST),'output_sha256':digest(DEST),'tables':records,'captions':captions,
 'references':selection,'original_preserved':True},ensure_ascii=False,indent=2),encoding='utf-8')
print('Saved revised application; four centered data tables and captions; five Chinese and three English references; original unchanged.')
