"""Create an editable execution book using bundled python-docx."""
from pathlib import Path
from docx import Document
from docx.shared import Pt,Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import re,json
ROOT=Path(__file__).resolve().parents[1];OUT=ROOT/'tmp/quality_book';OUT.mkdir(exist_ok=True,parents=True)
source=ROOT/'application/实施项目书.md';lines=source.read_text(encoding='utf-8').splitlines();doc=Document()
# The bundled default Title style contains a theme border. Remove inherited
# decoration before using a formal execution-book title.
for style in doc.styles:
    for border in style.element.xpath('.//w:pBdr'):border.getparent().remove(border)
section=doc.sections[0];section.page_width=Mm(210);section.page_height=Mm(297)
section.top_margin=Mm(24);section.bottom_margin=Mm(23);section.left_margin=Mm(25);section.right_margin=Mm(25);section.header_distance=Mm(12);section.footer_distance=Mm(12)
for name,size in [('Normal',11),('Title',22),('Heading 1',15),('Heading 2',12)]:
    st=doc.styles[name];st.font.name='Times New Roman';st.font.size=Pt(size);st.font.color.rgb=__import__('docx.shared',fromlist=['RGBColor']).RGBColor(0,0,0)
    st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体')
    for attr in ['asciiTheme','eastAsiaTheme','hAnsiTheme','cstheme']:
        st.element.get_or_add_rPr().rFonts.attrib.pop(qn('w:'+attr),None)
    pf=st.paragraph_format;pf.space_after=Pt(6);pf.space_before=Pt(0);pf.line_spacing=1.25;pf.keep_together=True
    if name in ('Heading 1','Heading 2'):st.font.bold=True;pf.keep_with_next=True;pf.keep_together=True;pf.space_before=Pt(8)
    if name=='Heading 1':pf.page_break_before=False
doc.styles['Title'].paragraph_format.space_after=Pt(24)
doc.styles['Title'].font.bold=True
header=section.header.paragraphs[0];header.text='南疆滴灌支管道故障诊断系统　实施执行稿';header.style=doc.styles['Normal']
for r in header.runs:r.font.size=Pt(9)
footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
footer.add_run('第 ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field);footer.add_run(' 页')
for r in footer.runs:r.font.size=Pt(9)
records={'source':source.relative_to(ROOT).as_posix(),'figures':[],'tables':[],'body_pt':11,'table_pt':10.5,'purpose':'autonomous_execution_book_not_official_taskbook_template'}
i=0
while i<len(lines):
    line=lines[i].strip();i+=1
    if not line:continue
    if line.startswith('# '):
        p=doc.add_paragraph(line[2:],'Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;continue
    if line.startswith('## '):
        p=doc.add_paragraph(line[3:],'Heading 1')
        if line.startswith('## 一 '):p.paragraph_format.page_break_before=True
        continue
    if line.startswith('### '):doc.add_paragraph(line[4:],'Heading 2');continue
    match=re.fullmatch(r'!\[(.+)\]\((.+)\)',line)
    if match:
        path=ROOT/'application'/match[2];width=390 if 'website_diagnosis_panel' in str(path) else 420
        p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_together=True
        p.add_run().add_picture(str(path),width=Pt(width));p.add_run('\n');r=p.add_run(match[1]);r.font.size=Pt(9.5)
        records['figures'].append({'caption':match[1],'path':path.relative_to(ROOT).as_posix(),'width_pt':width});continue
    if line.startswith('|'):
        raw=[line]
        while i<len(lines) and lines[i].strip().startswith('|'):raw.append(lines[i].strip());i+=1
        rows=[[x.strip() for x in r.strip('|').split('|')] for r in raw if not re.fullmatch(r'\|[ :\-|]+\|',r)]
        widths=([110,55,255] if rows[0][0]=='用途' else [105,150,165] if rows[0][0]=='时段' else [100,160,160] if rows[0][0]=='触发条件' else [100,170,150] if rows[0][0]=='正式材料要素' else [90,135,195])
        table=doc.add_table(rows=1,cols=3);table.autofit=False;table.style='Table Grid'
        borders=OxmlElement('w:tblBorders')
        for edge in ['top','left','bottom','right','insideH','insideV']:
            border=OxmlElement('w:'+edge);border.set(qn('w:val'),'single');border.set(qn('w:sz'),'4');border.set(qn('w:color'),'D9D9D9');borders.append(border)
        table._tbl.tblPr.append(borders)
        indent=OxmlElement('w:tblInd');indent.set(qn('w:w'),'120');indent.set(qn('w:type'),'dxa');table._tbl.tblPr.append(indent)
        for c,w in zip(table.columns,widths):c.width=Pt(w)
        for n,row in enumerate(rows):
            cells=table.rows[0].cells if n==0 else table.add_row().cells
            trPr=table.rows[n]._tr.get_or_add_trPr();trPr.append(OxmlElement('w:cantSplit'))
            if n==0:trPr.append(OxmlElement('w:tblHeader'))
            for c,value,w in zip(cells,row,widths):
                c.width=Pt(w);c.text=value;c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for p in c.paragraphs:
                    p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3);p.paragraph_format.line_spacing=1.15
                    p.paragraph_format.keep_with_next=n<len(rows)-1
                    for r in p.runs:r.font.size=Pt(10.5);r.bold=n==0
                    if rows[0][0]=='用途' and c==cells[1]:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                if n==0:
                    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'F0F0F0');c._tc.get_or_add_tcPr().append(shade)
        records['tables'].append({'header':rows[0],'rows':len(rows),'editable':True,'repeat_header':True,'widths_pt':widths})
        doc.add_paragraph().paragraph_format.space_after=Pt(1);continue
    p=doc.add_paragraph(line);p.alignment=WD_ALIGN_PARAGRAPH.LEFT if 'https://' in line else WD_ALIGN_PARAGRAPH.JUSTIFY;p.paragraph_format.widow_control=True
doc.core_properties.title='南疆滴灌支管道故障诊断系统实施项目书';doc.core_properties.subject='原型系统任务分解 实验协议 里程碑和验收';doc.core_properties.author='项目组 待确认'
path=OUT/'实施项目书.docx';doc.save(path)
(OUT/'layout_contract.json').write_text(json.dumps(records,ensure_ascii=False,indent=2),encoding='utf-8',newline='\n')
print('Editable execution book created, figures',len(records['figures']),'tables',len(records['tables']))
